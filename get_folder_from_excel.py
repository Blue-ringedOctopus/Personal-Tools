import os
import re
import shutil
import bisect
from collections import defaultdict
from datetime import datetime
from openpyxl import load_workbook

# ================== 配置路径（请按需修改） ==================
EXCEL_PATH = r"xx"
SOURCE_ROOT = r"xx"              # 源病历根目录
TARGET_ROOT = r"xx"  # 目标根目录
DEBUG = True  # 是否输出调试信息

# ================== 1. 读取 Excel ==================
def read_excel_dates(excel_path):
    """
    读取 B 列姓名、C 列日期，返回列表 [(姓名, YYYY-MM-DD), ...]
    自动去除姓名首尾空格，日期统一转换为 YYYY-MM-DD 字符串。
    """
    records = []
    wb = load_workbook(excel_path, data_only=True)
    ws = wb.active

    for row in ws.iter_rows(min_row=2, max_col=3, values_only=True):
        name = row[1]   # B 列
        date_val = row[2]  # C 列
        if not name or not date_val:
            continue

        # 姓名清洗
        name = str(name).strip()
        if not name:
            continue

        # 日期解析
        if isinstance(date_val, datetime):
            date_str = date_val.strftime("%Y-%m-%d")
        else:
            # 尝试多种常见格式
            date_str = None
            for fmt in ("%Y/%m/%d", "%Y-%m-%d", "%Y.%m.%d"):
                try:
                    dt = datetime.strptime(str(date_val), fmt)
                    date_str = dt.strftime("%Y-%m-%d")
                    break
                except ValueError:
                    continue
            if date_str is None:
                print(f"警告：无法解析日期 {date_val}，跳过")
                continue

        records.append((name, date_str))

    return records

# ================== 2. 文件系统扫描构建索引 ==================
def build_index_from_filesystem(source_root, debug=False):
    """
    通过扫描文件系统构建就诊索引。
    返回: dict { patient_name: [(date_str, full_folder_path), ...] }
    其中 date_str 为 'YYYY-MM-DD' 格式，full_folder_path 为日期文件夹的绝对路径。
    """
    index = defaultdict(list)
    date_pattern = re.compile(r'^\d{4}-\d{2}-\d{2}$')  # 匹配 YYYY-MM-DD

    # 遍历源根目录
    for root, dirs, files in os.walk(source_root):
        # 检查当前目录名是否为日期格式
        dir_name = os.path.basename(root)
        if date_pattern.match(dir_name):
            # 父目录名（应为患者姓名）
            parent_name = os.path.basename(os.path.dirname(root))
            # 排除父目录也为日期格式的情况（避免误匹配）
            if parent_name and not date_pattern.match(parent_name):
                # 检查路径中是否包含 "Aurum" 以确认是有效就诊记录
                if 'Aurum' in root:
                    patient_name = parent_name.strip()
                    if patient_name:
                        date_str = dir_name
                        index[patient_name].append((date_str, root))
                        if debug:
                            print(f"[索引] {patient_name} - {date_str} -> {root}")

    # 对每位患者的日期排序
    for patient in index:
        index[patient].sort(key=lambda x: x[0])  # 按日期字符串排序

    if debug:
        print(f"文件系统扫描完成，共索引到 {len(index)} 位患者。")
    return index

# ================== 3. 复制相邻日期文件夹 ==================
def copy_adjacent_records(records, source_index, target_root):
    """
    按患者分组，取 Excel 中该患者的最早和最晚日期，
    复制最早之前一次和最晚之后一次。
    返回: copied_count, skipped_count (文件夹级别)
    """
    # 按患者分组 Excel 日期
    patient_groups = defaultdict(list)
    for name, date in records:
        patient_groups[name].append(date)

    total_patients = len(patient_groups)
    stats = {
        'found_early': 0,
        'found_late': 0,
        'copied_early': 0,
        'copied_late': 0,
        'no_index': 0,
        'no_early_or_late': 0,
    }

    copied_count = 0
    skipped_count = 0
    copied_targets = set()  # 已复制过的目标路径，避免重复

    for name, excel_dates in patient_groups.items():
        if name not in source_index:
            print(f"⚠️ 未找到患者 '{name}' 的索引，跳过")
            stats['no_index'] += 1
            continue

        all_dates = source_index[name]  # list of (date_str, full_path)
        date_strs = [d[0] for d in all_dates]

        min_excel = min(excel_dates)
        max_excel = max(excel_dates)

        neighbors = []

        # ---- 寻找最早日期之前的一次 ----
        idx_min = bisect.bisect_left(date_strs, min_excel)
        if idx_min > 0:
            neighbors.append(all_dates[idx_min - 1])
            stats['found_early'] += 1

        # ---- 寻找最晚日期之后的一次 ----
        idx_max = bisect.bisect_left(date_strs, max_excel)
        if idx_max < len(date_strs):
            if date_strs[idx_max] == max_excel:
                if idx_max + 1 < len(date_strs):
                    neighbors.append(all_dates[idx_max + 1])
                    stats['found_late'] += 1
            else:
                neighbors.append(all_dates[idx_max])
                stats['found_late'] += 1

        if not neighbors:
            print(f"ℹ️ 患者 '{name}' 最早之前和最晚之后均不存在，跳过")
            stats['no_early_or_late'] += 1
            continue

        neighbor_dates = [d[0] for d in neighbors]
        print(f"处理患者 '{name}'，最早Excel日期 {min_excel}，最晚Excel日期 {max_excel}，将复制: {neighbor_dates}")

        for date_str, full_path in neighbors:
            dest_patient_dir = os.path.join(target_root, name)
            dest_dir = os.path.join(dest_patient_dir, date_str)

            if dest_dir in copied_targets:
                continue
            if os.path.exists(dest_dir):
                print(f"  目标已存在，跳过: {dest_dir}")
                skipped_count += 1
                copied_targets.add(dest_dir)
                continue

            os.makedirs(dest_patient_dir, exist_ok=True)
            try:
                shutil.copytree(full_path, dest_dir, ignore_dangling_symlinks=True)
                print(f"  ✅ 已复制: {full_path} -> {dest_dir}")
                copied_count += 1
                copied_targets.add(dest_dir)
                # 统计类型
                if date_str == neighbors[0][0] and stats['found_early']:
                    stats['copied_early'] += 1
                elif stats['found_late']:
                    stats['copied_late'] += 1
            except Exception as e:
                print(f"  ❌ 复制失败 {full_path} : {e}")

    # 输出统计报告
    print("\n" + "=" * 60)
    print("统计报告")
    print(f"总患者数（Excel去重后）   : {total_patients}")
    print(f"找到最早之前一次的患者数  : {stats['found_early']}")
    print(f"找到最晚之后一次的患者数  : {stats['found_late']}")
    print(f"成功复制最早之前的文件夹数: {stats['copied_early']}")
    print(f"成功复制最晚之后的文件夹数: {stats['copied_late']}")
    print(f"未找到索引的患者数        : {stats['no_index']}")
    print(f"无任何相邻日期的患者数    : {stats['no_early_or_late']}")
    print(f"成功复制的文件夹总数      : {copied_count}")
    print(f"因目标已存在而跳过的文件夹: {skipped_count}")
    print("=" * 60)

    return copied_count, skipped_count

# ================== 4. 主程序 ==================
def main():
    # 检查路径
    if not os.path.exists(EXCEL_PATH):
        print(f"❌ Excel文件不存在: {EXCEL_PATH}")
        return
    if not os.path.exists(SOURCE_ROOT):
        print(f"❌ 源目录不存在: {SOURCE_ROOT}")
        return

    print("正在读取 Excel...")
    records = read_excel_dates(EXCEL_PATH)
    print(f"共读取 {len(records)} 条记录。")

    print("正在扫描文件系统构建就诊索引...")
    source_index = build_index_from_filesystem(SOURCE_ROOT, debug=DEBUG)
    print(f"索引构建完成，共 {len(source_index)} 位患者有就诊记录。")

    print("开始复制相邻日期文件夹...")
    copy_adjacent_records(records, source_index, TARGET_ROOT)
    print("处理完成。")

if __name__ == "__main__":
    main()
