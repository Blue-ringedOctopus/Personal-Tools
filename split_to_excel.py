import re
import openpyxl
from openpyxl.styles import Font, Alignment

def parse_xiongbi_text(text):
    """
    解析校勘文本，支持三版本任意组合。
    方名标记位于“吴迁本：”之前，格式为“（xx方）吴迁本：”
    """
    pattern = r'(\d+)\.\s*([\s\S]*?)(?=\s*\d+\.\s*|$)'
    entries = re.findall(pattern, text)

    data = []
    for num, content in entries:
        def extract_version(version_name):
            # 匹配“版本名：”后面直到遇到下一个版本名或结束
            pattern = rf'{version_name}：([\s\S]*?)(?=\s*吴迁本：|\s*邓珍本：|\s*赵开美本：|$)'
            match = re.search(pattern, content)
            if not match:
                return '', '', ''
            raw = match.group(1).strip()
            # 分离附注（如“（对比吴迁本，共1处）”）
            comment_match = re.search(r'(\s*（[^）]*）)$', raw)
            if comment_match:
                comment = comment_match.group(1).strip()
                text_part = raw[:comment_match.start()].strip()
            else:
                text_part = raw
                comment = ''
            return text_part, comment

        # 先提取吴迁本原文，并尝试在“吴迁本：”前找方名标记
        # 使用更精确的方式：在content中搜索“（xx方）吴迁本：”
        fangming = ''
        wu_text = ''
        deng_text = ''
        zhao_text = ''
        wu_comm = ''
        deng_comm = ''
        zhao_comm = ''

        # 提取方名：匹配“（xx方）吴迁本：”
        fang_match = re.search(r'[（(]([^）)]+)[）)]\s*吴迁本：', content)
        if fang_match:
            fangming = fang_match.group(1).strip()
            # 然后提取吴迁本内容（从“吴迁本：”后面开始）
            wu_pattern = r'吴迁本：([\s\S]*?)(?=\s*邓珍本：|\s*赵开美本：|$)'
            wu_match = re.search(wu_pattern, content)
            if wu_match:
                wu_raw = wu_match.group(1).strip()
                wu_comm_match = re.search(r'(\s*（[^）]*）)$', wu_raw)
                if wu_comm_match:
                    wu_comm = wu_comm_match.group(1).strip()
                    wu_text = wu_raw[:wu_comm_match.start()].strip()
                else:
                    wu_text = wu_raw
        else:
            # 没有方名标记，正常提取吴迁本
            wu_pattern = r'吴迁本：([\s\S]*?)(?=\s*邓珍本：|\s*赵开美本：|$)'
            wu_match = re.search(wu_pattern, content)
            if wu_match:
                wu_raw = wu_match.group(1).strip()
                wu_comm_match = re.search(r'(\s*（[^）]*）)$', wu_raw)
                if wu_comm_match:
                    wu_comm = wu_comm_match.group(1).strip()
                    wu_text = wu_raw[:wu_comm_match.start()].strip()
                else:
                    wu_text = wu_raw

        # 提取邓珍本
        deng_pattern = r'邓珍本：([\s\S]*?)(?=\s*吴迁本：|\s*赵开美本：|$)'
        deng_match = re.search(deng_pattern, content)
        if deng_match:
            deng_raw = deng_match.group(1).strip()
            deng_comm_match = re.search(r'(\s*（[^）]*）)$', deng_raw)
            if deng_comm_match:
                deng_comm = deng_comm_match.group(1).strip()
                deng_text = deng_raw[:deng_comm_match.start()].strip()
            else:
                deng_text = deng_raw

        # 提取赵开美本
        zhao_pattern = r'赵开美本：([\s\S]*?)(?=\s*吴迁本：|\s*邓珍本：|$)'
        zhao_match = re.search(zhao_pattern, content)
        if zhao_match:
            zhao_raw = zhao_match.group(1).strip()
            zhao_comm_match = re.search(r'(\s*（[^）]*）)$', zhao_raw)
            if zhao_comm_match:
                zhao_comm = zhao_comm_match.group(1).strip()
                zhao_text = zhao_raw[:zhao_comm_match.start()].strip()
            else:
                zhao_text = zhao_raw

        data.append({
            '序号': num,
            '吴迁本原文': wu_text,
            '邓珍本原文': deng_text,
            '赵开美本原文': zhao_text,
            '吴本附注': wu_comm,
            '邓本附注': deng_comm,
            '赵本附注': zhao_comm,
            '备注': fangming,
        })

    return data


def fill_from_template(template_path, output_path, parsed_data):
    wb = openpyxl.load_workbook(template_path)
    ws = wb['校勘分类']

    col_index = {
        '序号': 1,
        '吴迁本': 2,
        '邓珍本': 3,
        '赵开美本': 4,
        '校勘对象': 5,
        '致误原因': 6,
        '判断依据': 7,
        '备注': 8,
        '邓本附注': 9,
        '赵本附注': 10,
    }

    for row_idx, row_data in enumerate(parsed_data, start=2):
        ws.cell(row=row_idx, column=col_index['序号']).value = row_data['序号']
        ws.cell(row=row_idx, column=col_index['吴迁本']).value = row_data['吴迁本原文']
        ws.cell(row=row_idx, column=col_index['邓珍本']).value = row_data['邓珍本原文']
        ws.cell(row=row_idx, column=col_index['赵开美本']).value = row_data['赵开美本原文']
        ws.cell(row=row_idx, column=col_index['备注']).value = row_data['备注']
        ws.cell(row=row_idx, column=col_index['邓本附注']).value = row_data['邓本附注']
        ws.cell(row=row_idx, column=col_index['赵本附注']).value = row_data['赵本附注']

    font_style = Font(name='宋体', size=11)
    align_style = Alignment(wrap_text=True, vertical='top')
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=10):
        for cell in row:
            cell.font = font_style
            cell.alignment = align_style

    ws.column_dimensions['A'].width = 8
    ws.column_dimensions['B'].width = 60
    ws.column_dimensions['C'].width = 60
    ws.column_dimensions['D'].width = 60
    ws.column_dimensions['E'].width = 15
    ws.column_dimensions['F'].width = 15
    ws.column_dimensions['G'].width = 40
    ws.column_dimensions['H'].width = 25
    ws.column_dimensions['I'].width = 30
    ws.column_dimensions['J'].width = 30

    wb.save(output_path)
    print(f"✅ 成功导出！共处理 {len(parsed_data)} 条校勘条目。")
    print(f"📁 文件保存为：{output_path}")


if __name__ == "__main__":
    txt_path = "G:\金匮要略\惊悸吐衄第十六\新建文本文档.txt"
    template_path = "G:\金匮要略\空白表格.xlsx"
    output_path = "G:\金匮要略\惊悸吐衄第十六\惊悸吐衄第十六_校勘分类.xlsx"

    try:
        with open(txt_path, 'r', encoding='utf-8') as f:
            text = f.read()
    except FileNotFoundError:
        print(f"❌ 找不到文件：{txt_path}")
        exit()

    parsed = parse_xiongbi_text(text)
    fill_from_template(template_path, output_path, parsed)